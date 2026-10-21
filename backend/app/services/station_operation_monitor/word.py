from __future__ import annotations

from copy import deepcopy
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import _Row

from app.utils.path_config import resolve_agent_path
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .core import local_time


FIRST_PAGE_NOTE = ('1 东莞市的南城元岭站、莞城梨川站与东城主山站，在“十四五”期间均为国控站点。'
                   '根据《“十五五”国家城市环境空气质量监测点位设置方案》，该三个站点已完成交接，'
                   '数据传输正常，现纳入省控站点监控。')


def apply_first_page_note(document):
    """Approved fixed first-page footer; normal page numbering remains intact."""
    section = document.sections[0]
    primary_footer = section.footer
    primary_header = section.header
    section.different_first_page_header_footer = True
    for target, source in [(section.first_page_footer, primary_footer),
                           (section.first_page_header, primary_header)]:
        target.is_linked_to_previous = False
        for child in list(target._element):
            target._element.remove(child)
        for child in source._element:
            target._element.append(deepcopy(child))
    footer = section.first_page_footer
    note = footer.paragraphs[0].insert_paragraph_before(FIRST_PAGE_NOTE)
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    layout = note.paragraph_format
    layout.left_indent = layout.right_indent = layout.first_line_indent = Pt(0)
    layout.space_before = Pt(0)
    layout.space_after = Pt(6)
    layout.line_spacing = Pt(14)
    layout.keep_together = True
    for run in note.runs:
        run.font.name = '仿宋_GB2312'
        run.font.size = Pt(11)
        run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '仿宋_GB2312')


def stamp(value):
    return local_time(value).strftime('%m月%d日%H时')


def write_word(data, config, path):
    if config.get("template_path"):
        return write_template_word(data, config, path)
    document = Document()
    style = document.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    document.add_heading(config.get('report_title', '空气质量国控和省控站点运行监控日报'), 0)
    document.add_paragraph(local_time(data['window_end']).strftime('（%Y年%m月%d日）'))
    document.add_paragraph(config.get('organization', ''))
    document.add_paragraph(f"统计时段：{stamp(data['window_start'])}至{stamp(data['window_end'])}。")
    document.add_paragraph('通过站点小时数据接口检查、核对，整理连续缺数超过4小时的数据缺失记录，同时跟进前日及更早未结记录。')
    if not data['report_ready']:
        document.add_paragraph('本次取数不完整，报告为待核实稿；不能据此认定所有站点运行正常。')
    for kind in ('国控', '省控'):
        entries = [entry for entry in data['entries'] if entry['type'] == kind]
        document.add_heading(f"{kind}站点（本次监控范围{data['station_counts'][kind]}个）：", 1)
        if not entries:
            document.add_paragraph('本次监控范围内未发现已观察的小时数据缺失记录。')
        for entry in entries:
            restored = '截至统计截止小时数据已返回，设备恢复时间待核实' if entry['recovered'] else '截至统计截止时未恢复，正在持续跟进'
            filled = '数据已回补' if entry['backfill_status'] == '已回补' else '数据'+entry['backfill_status']
            if entry.get('unrecoverable_confirmation'):
                filled += '（'+entry['unrecoverable_confirmation']['reason']+'）'
            prefix = '前日及更早情况跟进：' if entry['carry_over'] else ''
            document.add_paragraph(f"{prefix}{entry['city']}{entry['name']}站：{stamp(entry['offline_start'])}至{stamp(entry['offline_end'])}，"
                                   f"累计影响数据{entry['affected_hours']}小时（本统计窗口{entry['window_affected_hours']}小时），{restored}，{filled}。")
    document.add_paragraph('编制、校核及期号：待人工确认。')
    document.add_heading('附件：站点离线记录表（按连续缺数超过4小时判定）', 1)
    headers = ['序号','城市','站点','离线时间','是否已恢复','影响数据时长（小时）','是否影响数据有效性','排查结果','备注']
    for kind in ('国控', '省控'):
        document.add_heading(f'{kind}站点离线记录表', 2)
        table = document.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        entries = [entry for entry in data['entries'] if entry['type'] == kind]
        for i, entry in enumerate(entries, 1):
            note = '数据已回补' if entry['backfill_status']=='已回补' else entry['backfill_status']
            if entry.get('unrecoverable_confirmation'):
                note += '；'+entry['unrecoverable_confirmation']['reason']+'；依据：'+entry['unrecoverable_confirmation']['evidence']
            if entry['backfill_sources']:
                note += '；' + '、'.join(entry['backfill_sources'])
            if entry['carry_over']:
                note += '；历史未结跟进'
            values = [str(i),entry['city'],entry['name'],f"{stamp(entry['offline_start'])}—{stamp(entry['offline_end'])}",
                      '是' if entry['recovered'] else '否',str(entry['affected_hours']), '待核实',
                      entry.get('issue_reason', '原因待核实'),note]
            if entry.get('feedback'):
                values[-1] += '；' + entry['feedback']['feedback_by'] + '反馈，' + stamp(entry['feedback']['feedback_at']) + '起关闭跟进'
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        if not entries:
            for cell in table.add_row().cells:
                cell.text = '/'
    if data['unverified_stations']:
        document.add_heading('无记录候选站点（在运状态待核实）', 1)
        document.add_paragraph('以下候选站点未取得数据，尚未纳入已确认离线统计：' + '；'.join(
            f"{s['city']}{s['name']}（{s['code']}，{s['type']}）" for s in data['unverified_stations']))
    if data['collection_errors']:
        document.add_heading('取数异常', 1)
        for error in data['collection_errors']:
            document.add_paragraph('；'.join(str(v) for v in error.values()))
    document.add_heading('数据来源及说明', 1)
    document.add_paragraph('国控：'+'、'.join(data['national_sources'])+'。省控：'+data['provincial_source']+'。')
    for limitation in data['limitations']:
        document.add_paragraph(limitation)
    document.save(path)
    return path


def replace_text(paragraph, text):
    """Replace text while retaining the template paragraph and first run format."""
    run = next((r for r in paragraph.runs if r.text), None)
    properties = deepcopy(run._r.rPr) if run is not None and run._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag != qn('w:pPr'):
            paragraph._p.remove(child)
    new_run = paragraph.add_run(text)
    if properties is not None:
        new_run._r.insert(0, properties)



def replace_spans(paragraph, spans):
    properties = [deepcopy(r._r.rPr) for r in paragraph.runs]
    replace_text(paragraph, '')
    for text, index in spans:
        run = paragraph.add_run(text)
        if properties[index] is not None:
            run._r.insert(0, deepcopy(properties[index]))

def station_label(entry):
    city, name = entry['city'], entry['name']
    return (city if city.endswith('市') else city + '市') + (name if name.endswith('站') else name + '站')


def incident_text(entry):
    """Use the approved daily-report wording; retain technical evidence in JSON."""
    reason = entry.get('issue_reason')
    cause = f"经核实为{reason}，" if reason and reason != '/' else ''
    recovery = '现已恢复正常' if entry['recovered'] else '正在持续跟进'
    backfill = ''
    if entry['backfill_status'] == '已回补':
        backfill = '，数据已回补'
    elif entry['backfill_status'] == '无法回补':
        backfill = '，数据无法回补'
    elif entry['recovered']:
        backfill = '，数据待回补'
    return (f"{station_label(entry)}：{cause}累计影响数据时长{entry['affected_hours']}小时，"
            f"{recovery}{backfill}。")


def incident_table_values(entry, index):
    start, end = local_time(entry['offline_start']), local_time(entry['offline_end'])
    left = f'{start.month}月{start.day}日{start.hour}时'
    if (start.year, start.month, start.day) == (end.year, end.month, end.day):
        right = f'{end.hour}时'
    elif (start.year, start.month) == (end.year, end.month):
        right = f'{end.day}日{end.hour}时'
    else:
        right = f'{end.month}月{end.day}日{end.hour}时'
    status = entry['backfill_status']
    note = ('数据已回补' if status == '已回补' else '数据无法回补' if status == '无法回补'
            else '数据待回补' if entry['recovered'] else '持续跟进')
    validity = entry.get('data_validity_impact')
    return [str(index), entry['city'], entry['name'], left + '-' + right,
            '是' if entry['recovered'] else '否', str(entry['affected_hours']),
            validity if validity in ('是', '否') else '/', entry.get('issue_reason') or '/', note]


def section_summary(kind, entries):
    count = len({(entry['city'], entry['name']) for entry in entries})
    if count:
        return f'过去24小时内，有{count}个{kind}站点出现过离线超4小时的情况，经核实，情况如下：'
    return f'过去24小时内，无{kind}站点出现过离线超4小时的情况。'


def overview_summary(data):
    end = local_time(data['window_end'])
    prefix = (f'通过广东省空气质量监测管理与发布系统检查、核对，截至{end.month}月{end.day}日{end.hour}时，'
              '在过去24小时内，')
    if data.get('report_ready') is False:
        return prefix + '本次取数未完成，运行情况待核实。'
    totals = {kind: len({(e['city'], e['name']) for e in data['entries'] if e['type'] == kind})
              for kind in ('国控', '省控')}
    affected = '和'.join(f'{totals[kind]}个{kind}站点' for kind in ('国控', '省控') if totals[kind])
    if affected:
        return prefix + f'有{affected}出现过离线超4小时的情况，其余国控和省控站点运行正常。'
    return prefix + '国控和省控站点运行正常。'


def write_template_word(data, config, path):
    template_path = resolve_agent_path(config['template_path'])
    digest = hashlib.sha256(template_path.read_bytes()).hexdigest()
    if config.get('template_sha256') and config['template_sha256'] != digest:
        raise ValueError('日报模板版本校验失败')
    document = Document(template_path)
    apply_first_page_note(document)
    paragraphs = document.paragraphs
    # This layout adapter deliberately validates the supplied Guangdong report structure.
    if (len(paragraphs) < 20 or len(document.tables) != 2
            or paragraphs[1].text != '运行监控日报'
            or paragraphs[6].text[:4] != '国控站点'
            or paragraphs[10].text[:4] != '省控站点'
            or any(len(t.columns) != 9 for t in document.tables)):
        raise ValueError('日报模板结构不匹配，请核对模板版本')
    end = local_time(data['window_end'])
    replace_text(paragraphs[2], f'（{end.year}年{end.month}月{end.day}日）')
    replace_text(paragraphs[3], config.get('organization', '广东省生态环境监测中心') +
                 '       ' + config.get('report_issue', f'{end.year}年期号待确认'))
    counts = data['station_counts']
    replace_text(paragraphs[5], overview_summary(data))
    replace_text(paragraphs[6], f"国控站点（{counts['国控']}个）：")
    replace_text(paragraphs[10], f"省控站点（{counts['省控']}个）：")
    prototype = deepcopy(paragraphs[8]._p)
    for i in (8, 9):
        paragraphs[i]._p.getparent().remove(paragraphs[i]._p)
    for kind, anchor in [('国控', paragraphs[7]), ('省控', paragraphs[11])]:
        entries = [e for e in data['entries'] if e['type'] == kind]
        replace_text(anchor, section_summary(kind, entries))
        for entry in entries:
            element = deepcopy(prototype)
            anchor._p.addnext(element)
            anchor = Paragraph(element, anchor._parent)
            text = incident_text(entry)
            label, detail = text.split('：', 1)
            replace_spans(anchor, [(label + '：', 0), (detail, 1)])
    populated_cell_formats = [deepcopy(c.paragraphs[0].runs[0]._r.rPr)
                              for c in document.tables[0].rows[1].cells]
    for kind, table in zip(('国控', '省控'), document.tables):
        row_template = deepcopy(table.rows[1]._tr)
        for row in list(table.rows)[1:]:
            table._tbl.remove(row._tr)
        entries = [e for e in data['entries'] if e['type'] == kind]
        rows = entries or [None]
        for index, entry in enumerate(rows, 1):
            row_element = deepcopy(row_template)
            table._tbl.append(row_element)
            row = _Row(row_element, table)
            if entry is None:
                values = ['/'] * 9
            else:
                values = incident_table_values(entry, index)
            for column, (cell, value) in enumerate(zip(row.cells, values)):
                replace_text(cell.paragraphs[0], value)
                # The original empty provincial time cell inherits body size.
                # Populated rows in both additional samples use the national row's 9pt format.
                run = cell.paragraphs[0].runs[0]
                fallback = populated_cell_formats[column]
                if entry is not None and run.font.size is None and fallback is not None:
                    if run._r.rPr is not None:
                        run._r.remove(run._r.rPr)
                    run._r.insert(0, deepcopy(fallback))
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
    document.save(path)
    validation = validate_template_format(template_path, path, data=data)
    validation['template_sha256'] = digest
    Path(path).with_suffix('.format.json').write_text(
        json.dumps(validation, ensure_ascii=False, indent=2), encoding='utf-8')
    return path


def validate_template_format(template_path, output_path, *, data=None):
    """Validate stable layout invariants, allowing variable body content and row counts."""
    template = Document(template_path)
    apply_first_page_note(template)
    output = Document(output_path)
    checks = {}
    checks['sections'] = [s._sectPr.xml for s in template.sections] == [s._sectPr.xml for s in output.sections]
    checks['headers_footers'] = len(template.sections) == len(output.sections) and all(
        a.header._element.xml == b.header._element.xml and a.footer._element.xml == b.footer._element.xml
        and a.first_page_header._element.xml == b.first_page_header._element.xml
        and a.first_page_footer._element.xml == b.first_page_footer._element.xml
        for a, b in zip(template.sections, output.sections))
    note_paragraphs = [p for p in output.sections[0].first_page_footer.paragraphs if p.text == FIRST_PAGE_NOTE]
    checks['first_page_note'] = (output.sections[0].different_first_page_header_footer
        and len(note_paragraphs) == 1 and bool(note_paragraphs[0].runs)
        and all(r.font.size == Pt(11) for r in note_paragraphs[0].runs)
        and not any(FIRST_PAGE_NOTE in p.text for section in output.sections for p in section.footer.paragraphs)
        and not any(section.different_first_page_header_footer for section in output.sections[1:]))
    checks['titles'] = len(output.paragraphs) >= 5 and all(
        template.paragraphs[i]._p.xml == output.paragraphs[i]._p.xml for i in (0, 1))
    checks['date_issuer_styles'] = len(output.paragraphs) >= 5 and all(
        template.paragraphs[i]._p.pPr.xml == output.paragraphs[i]._p.pPr.xml
        and template.paragraphs[i].runs[0]._r.rPr.xml == output.paragraphs[i].runs[0]._r.rPr.xml
        for i in (2, 3))
    checks['divider'] = len(output.paragraphs) >= 5 and template.paragraphs[4]._p.xml == output.paragraphs[4]._p.xml
    checks['tables'] = len(template.tables) == len(output.tables) == 2 and all(
        a._tbl.tblPr.xml == b._tbl.tblPr.xml and a._tbl.tblGrid.xml == b._tbl.tblGrid.xml
        and a.rows[0]._tr.xml == b.rows[0]._tr.xml
        for a, b in zip(template.tables, output.tables))
    checks['table_titles'] = all(
        any(p._p.xml == template.paragraphs[i]._p.xml for p in output.paragraphs) for i in (16, 18))
    if data is not None:
        # Only these template paragraphs and the two table data areas are mutable.
        editable = {2, 3, 5, 6, 7, 8, 9, 10, 11}
        fixed_template = [p._p.xml for i, p in enumerate(template.paragraphs) if i not in editable]
        incident_count = len(data['entries'])
        n_national = sum(e['type'] == '国控' for e in data['entries'])
        editable_output = {2, 3, 5, 6, 7, 8+n_national, 9+n_national}
        editable_output.update(range(8, 8+n_national))
        editable_output.update(range(10+n_national, 10+incident_count))
        fixed_output = [p._p.xml for i, p in enumerate(output.paragraphs) if i not in editable_output]
        checks['fixed_content'] = fixed_template == fixed_output and len(output.paragraphs) == 18 + incident_count
        checks['fixed_signatures_recipients'] = all(
            output.paragraphs[index+incident_count]._p.xml == template.paragraphs[source]._p.xml
            for source, index in [(12, 10), (13, 11)])
        checks['overview_wording'] = output.paragraphs[5].text == overview_summary(data)
        checks['section_wording'] = all(
            output.paragraphs[index].text == section_summary(kind, [e for e in data['entries'] if e['type'] == kind])
            for kind, index in [('国控', 7), ('省控', 9+n_national)])
        checks['event_rows'] = all(len(t.rows) == 1 + max(1, sum(e['type'] == kind for e in data['entries']))
                                   for kind, t in zip(('国控', '省控'), output.tables))
    failures = [name for name, passed in checks.items() if not passed]
    if failures:
        raise ValueError('日报格式校验失败：' + '、'.join(failures))
    return {'passed': True, 'checks': checks}


def render_pdf_preview(word_path):
    """Render the exact generated Word using an isolated LibreOffice profile."""
    from tempfile import TemporaryDirectory
    from app.tools.office.soffice import run_soffice

    word_path = Path(word_path).resolve()
    with TemporaryDirectory(prefix='station-report-office-') as profile:
        result = run_soffice([
            '-env:UserInstallation=' + Path(profile).as_uri(), '--headless',
            '--convert-to', 'pdf', '--outdir', str(word_path.parent), str(word_path),
        ], capture_output=True, timeout=120)
    pdf_path = word_path.with_suffix('.pdf')
    if result.returncode != 0 or not pdf_path.is_file() or not pdf_path.read_bytes().startswith(b'%PDF-'):
        raise RuntimeError('日报 PDF 预览生成失败')
    return pdf_path
