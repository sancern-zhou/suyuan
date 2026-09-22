from types import SimpleNamespace

import pytest

from app.tools.utility.execute_python_tool import ExecutePythonTool


def context_for(directory):
    return SimpleNamespace(
        data_manager=SimpleNamespace(memory=SimpleNamespace(
            session=SimpleNamespace(data_dir=directory))),
        available_file_paths=[], authorized_input_paths=[],
    )


@pytest.mark.asyncio
async def test_declared_inputs_support_dynamic_paths_and_excel(tmp_path):
    sources = [tmp_path / 'one.json', tmp_path / 'two.json']
    for index, source in enumerate(sources):
        source.write_text('[{"value": %d}]' % index)
    tool = ExecutePythonTool()
    context = context_for(tmp_path)
    result = await tool.execute(context=context, input_files=list(map(str, sources)), code="""
from pathlib import Path
from openpyxl import Workbook, load_workbook
rows = []
for source in input_files:
    path = Path(source)
    rows.extend(load_data(str(path.parent / path.name)))
workbook = Workbook()
for row in rows:
    workbook.active.append([row['value']])
destination = artifact_path('declared_inputs.xlsx')
workbook.save(destination)
assert list(load_workbook(destination).active.values) == [(0,), (1,)]
print('EXCEL_OK')
""")
    assert result['success'] is True, result
    assert 'EXCEL_OK' in result['data']['output']
    assert result.get('resources'), result
    assert sources[0].read_text() == '[{"value": 0}]'
    denied = await tool.execute(context=context, input_files=[],
                                code=f"load_data({str(sources[0])!r})")
    assert denied['success'] is False, denied


@pytest.mark.asyncio
@pytest.mark.parametrize('kind', ['outside', 'symlink', 'missing', 'directory', 'invalid'])
async def test_invalid_inputs_fail_before_execution(tmp_path, kind):
    directory = tmp_path / 'session'
    directory.mkdir()
    outside = tmp_path / 'outside.json'
    outside.write_text('[]')
    source = directory / 'missing.json'
    if kind == 'outside':
        source = outside
    elif kind == 'symlink':
        source.symlink_to(outside)
    elif kind == 'directory':
        source = directory
    inputs = [str(source)] if kind != 'invalid' else 'not-an-array'
    result = await ExecutePythonTool().execute(
        context=context_for(directory), input_files=inputs,
        code="raise AssertionError('must not execute')")
    assert result['error_code'] == 'INVALID_INPUT_FILES', result


def test_authorized_external_input_and_schema(tmp_path):
    source = tmp_path / 'upload.json'
    source.write_text('[]')
    context = context_for(tmp_path / 'session')
    context.authorized_input_paths = [str(source)]
    tool = ExecutePythonTool()
    assert tool._validate_input_files([str(source)], context) == [str(source)]
    schema = tool.get_function_schema()
    assert 'input_files' in schema['parameters']['properties']
    assert '独立环境' in schema['description']
