from datetime import datetime,timedelta

import pytest

from app.services.station_operation_monitor.core import MonitorStore, hours, local_time, normalize, parse_value
from app.services.station_operation_monitor.runtime import collect, generate_report


BASE=local_time('2026-10-06T10:00:00')
ROSTER=[{'code':'A','name':'一站','city':'测试市','type':'国控','confirmed':True,'active':True},
        {'code':'B','name':'二站','city':'测试市','type':'省控','confirmed':True,'active':True},
        {'code':'OLD','name':'旧站','city':'测试市','type':'国控'}]


def rows(start,end,codes=('A','B'),missing=()):
    return [{'code':code,'timePoint':hour,'sO2':'7(H)','nO2':'0','pM10':'11(Re)',
             'co':'0.4','o3':'30','pM2_5':'8'} for code in codes for hour in hours(start,end)
            if (code,hour) not in missing]


def ingest(store,start,end,records,source='primary',observed=None,**kwargs):
    store.ingest(source=source,city='测试市',start=start,end=end,records=records,
                 observed_at=observed or end+timedelta(minutes=30),**kwargs)


def report(store,end):
    return store.report(end,national_sources=['primary','secondary'],provincial_source='primary')


def test_hour_grid_and_quality_flags():
    assert len(list(hours(BASE,BASE+timedelta(days=1))))==24
    assert parse_value('7(H)')==(7,'H')
    assert parse_value('0')==(0,None)
    assert parse_value('-99')[0] is None
    assert parse_value('—')[0] is None
    r=normalize(rows(BASE,BASE+timedelta(hours=1))[0])
    assert r['present'] and not r['quality_complete']
    assert r['marks']['sO2']=='H' and r['marks']['pM10']=='Re'


def test_failure_does_not_create_offline_and_old_stations_not_active(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    ingest(store,BASE,end,rows(BASE,end))
    ingest(store,BASE,end,[],success=False,detail='timeout')
    r=report(store,end)
    assert not r['entries']
    assert not r['report_ready']
    assert [s['code'] for s in r['unverified_stations']]==['OLD']


def test_recovered_gap_leaves_tracking_without_backfill(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    missing_hours=list(hours(BASE,BASE+timedelta(hours=5)))
    absent={('A',h) for h in missing_hours}
    end=BASE+timedelta(days=1)
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    entry=report(store,end)['entries'][0]
    assert entry['affected_hours']==5 and entry['recovered']
    assert entry['backfill_status']=='待回补'
    # Recovery alone closes tracking; pending backfill must not extend it.
    next_end=end+timedelta(days=2)
    assert report(store,next_end)['entries']==[]
    # A late backfill must not resurrect the closed incident either.
    ingest(store,BASE,end,rows(BASE,end,codes=('A',)),source='secondary',station_types=['国控'],
           observed=next_end+timedelta(minutes=30))
    assert report(store,next_end)['entries']==[]


def test_either_source_must_fill_entire_gap_not_mixed_hours(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    absent={('A',h) for h in list(hours(BASE,BASE+timedelta(hours=5)))}
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    # Fill only three of five through second source: still pending.
    ingest(store,BASE,BASE+timedelta(hours=3),rows(BASE,BASE+timedelta(hours=3),codes=('A',)),
           source='secondary',station_types=['国控'],observed=end+timedelta(minutes=31))
    assert report(store,end)['entries'][0]['backfill_status']=='待回补'
    ingest(store,BASE,end,rows(BASE,end,codes=('A',)),source='secondary',station_types=['国控'],
           observed=end+timedelta(minutes=32))
    entry=report(store,end)['entries'][0]
    assert entry['backfill_status']=='已回补' and entry['backfill_sources']==['secondary']
    # Originally observed evidence survives the fill.
    assert entry['affected_hours']==5


def test_stale_recheck_cannot_erase_later_backfill(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    ingest(store,BASE,end,[],observed=end+timedelta(minutes=20))
    ingest(store,BASE,end,rows(BASE,end),observed=end+timedelta(minutes=30))
    ingest(store,BASE,end,[],observed=end+timedelta(minutes=25))
    entry=report(store,end)['entries'][0]
    assert entry['backfill_status']=='已回补' and entry['recovered']


def test_only_gaps_over_four_hours_are_recorded(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    missing={('A',h) for h in hours(BASE,BASE+timedelta(hours=4))}
    missing.update(('B',h) for h in hours(BASE,BASE+timedelta(hours=5)))
    ingest(store,BASE,end,rows(BASE,end,missing=missing))
    # 4-hour gap stays below the threshold; 5-hour gap becomes an incident.
    entries=report(store,end)['entries']
    assert [e['code'] for e in entries]==['B'] and entries[0]['affected_hours']==5
    assert entries[0]['evidence_kind'].startswith('连续小时数据缺失超过4小时')
    assert any('连续缺数超过4小时' in item for item in report(store,end)['limitations'])


class FakeSource:
    calls=[]
    def __init__(self,config):pass
    def fetch(self,codes,start,end):
        self.calls.append((codes,start,end))
        return rows(start,end,codes=[c for c in codes if c!='OLD'])


def config():
    return {'stations':ROSTER,'sources':[{'name':'primary'}],
            'national_sources':['primary'],'provincial_source':'primary','bootstrap_hours':48}


def test_collector_slicing_raw_preservation_and_word(tmp_path):
    now=BASE+timedelta(days=1,minutes=30)
    result=collect(config(),now=now,root=tmp_path,source_factory=FakeSource)
    assert result['success']
    assert len(list((tmp_path/'snapshots').rglob('*.json.gz')))>=2
    r=generate_report(config(),now=now,root=tmp_path,refresh=False)
    assert r['report_ready'] and r['unverified_station_count']==1
    from docx import Document
    content='\n'.join(p.text for p in Document(r['word_path']).paragraphs)
    assert '独立发布源未接入' in content
    assert '无记录候选站点' in content


def test_arrival_grace_and_future_window_rejected(tmp_path):
    with pytest.raises(ValueError):
        collect(config(),now=BASE+timedelta(minutes=10),end=BASE,root=tmp_path,source_factory=FakeSource)
    with pytest.raises(ValueError):
        generate_report(config(),now=BASE,root=tmp_path,refresh=False)


def test_partial_return_does_not_recover_missing_parameters(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    absent={('B',h) for h in hours(BASE,BASE+timedelta(hours=5))}
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    partial=rows(BASE,end)
    for record in partial:
        if record['code']=='B':record['co']='—'
    ingest(store,BASE,end,partial,observed=end+timedelta(minutes=40))
    entry=report(store,end)['entries'][0]
    assert not entry['recovered'] and entry['backfill_status']=='待回补'


def test_malformed_response_not_persisted_as_station_outage(tmp_path):
    class Malformed(FakeSource):
        def fetch(self,codes,start,end):
            return [{'code':'A','timePoint':'not-a-time'}]
    result=collect(config(),now=BASE+timedelta(days=1,minutes=30),root=tmp_path,source_factory=Malformed)
    assert not result['success']
    store=MonitorStore(tmp_path/'monitor.sqlite3')
    assert not report(store,BASE+timedelta(days=1))['entries']


def test_http_empty_and_truncated_response_are_not_accepted():
    from app.services.station_operation_monitor.runtime import StationHourHTTPSource
    class Response:
        status_code=200
        def raise_for_status(self):pass
        def json(self):return {'success':True,'result':[]}
    class Session:
        def post(self,*args,**kwargs):return Response()
    source=StationHourHTTPSource({'base_url':'https://example.invalid','hour_endpoint':'/hours'})
    source.token='test';source.session=Session()
    with pytest.raises(RuntimeError,match='empty'):
        source.fetch(['A'],BASE,BASE+timedelta(hours=1))


def test_report_no_observations_is_incomplete(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    r=report(store,BASE+timedelta(days=1))
    assert not r['report_ready'] and not r['entries']


@pytest.mark.asyncio
async def test_fetcher_propagates_collection_failure(monkeypatch):
    from app.fetchers import station_operation_monitor as module
    monkeypatch.setattr(module,'load_config',lambda:{'collection_cron':'25 * * * *'})
    monkeypatch.setattr(module,'collect',lambda:{'success':False,'failed_slices':[{}]})
    fetcher=module.StationOperationMonitorFetcher()
    assert fetcher.schedule=='25 * * * *'
    with pytest.raises(RuntimeError):
        await fetcher.fetch_and_store()


def test_unrecoverable_requires_attributable_confirmation(tmp_path):
    store=MonitorStore(tmp_path/'monitor.sqlite3');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    absent={('B',h) for h in hours(BASE,BASE+timedelta(hours=5))}
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    cfg=config()
    confirmation={'code':'B','start':BASE.isoformat(),'end':(BASE+timedelta(hours=5)).isoformat(),
                  'reason':'原始记录丢失','confirmed_by':'运维核实','confirmed_at':end.isoformat()}
    cfg['unrecoverable_confirmations']=[confirmation]
    r=generate_report(cfg,now=end+timedelta(minutes=30),root=tmp_path,refresh=False)
    import json
    assert json.load(open(r['report_input_path']))['entries'][0]['backfill_status']=='待回补'
    confirmation['evidence']='工单TEST-001的核实记录'
    r=generate_report(cfg,now=end+timedelta(minutes=30),root=tmp_path,refresh=False)
    assert json.load(open(r['report_input_path']))['entries'][0]['backfill_status']=='无法回补'


def test_continuously_offline_events_remain_after_three_report_days(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    gap_end=BASE+timedelta(hours=5)
    absent={('A',h) for h in hours(BASE,gap_end)}
    # No data returns after the gap, so the incident must remain beyond day three.
    ingest(store,BASE,gap_end,rows(BASE,gap_end,codes=('A',),missing=absent),
           observed=gap_end+timedelta(minutes=30),station_types=['国控'])
    listed=[[e['code'] for e in report(store,BASE+timedelta(days=day))['entries']]
            for day in (1,2,3,4,10)]
    assert listed==[['A']]*5


def test_record_feedback_closes_tracking_and_fills_reason(tmp_path):
    from app.services.station_operation_monitor.feedback import record_feedback
    store=MonitorStore(tmp_path/'monitor.sqlite3');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    absent={('B',h) for h in hours(BASE,BASE+timedelta(hours=5))}
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    cfg=config()
    feedback=record_feedback(cfg,station='B',reason='光缆中断',unrecoverable=True,
                             feedback_by='李四',feedback_at=end-timedelta(hours=1),root=tmp_path)
    assert feedback['record']['code']=='B' and feedback['matched_events']
    result=generate_report(cfg,now=end+timedelta(minutes=40),root=tmp_path,refresh=False)
    import json
    entries=json.load(open(result['report_input_path']))['entries']
    assert entries and entries[0]['issue_reason']=='光缆中断'
    assert entries[0]['backfill_status']=='无法回补'
    from docx import Document
    table_text='\n'.join(cell.text for table in Document(result['word_path']).tables
                         for row in table.rows for cell in row.cells)
    assert '光缆中断' in table_text and '李四' in table_text
    # Feedback recorded before the report day closes tracking entirely.
    result2=generate_report(cfg,now=end+timedelta(days=1,minutes=40),root=tmp_path,refresh=False)
    assert not json.load(open(result2['report_input_path']))['entries']


def test_unrecoverable_feedback_stops_pending_refetch(tmp_path):
    from app.services.station_operation_monitor.feedback import record_feedback
    from app.services.station_operation_monitor.runtime import collect

    class Spy:
        calls=[]
        def __init__(self,config):pass
        def fetch(self,codes,start,end):
            self.calls.append((tuple(codes),start,end))
            return rows(start,end,codes=[c for c in codes if c!='OLD'])

    def prepare(root):
        store=MonitorStore(root/'monitor.sqlite3');store.seed(ROSTER)
        gap_start=BASE-timedelta(days=2)
        absent={('B',h) for h in hours(gap_start+timedelta(hours=2),gap_start+timedelta(hours=4))}
        # Window ends inside the outage so no later complete hour exists and the
        # gap stays pending (ongoing offline).
        ingest(store,gap_start,gap_start+timedelta(hours=3),
               rows(gap_start,gap_start+timedelta(hours=3),missing=absent),
               observed=gap_start+timedelta(hours=3,minutes=30))

    control=tmp_path/'control';control.mkdir();prepare(control)
    Spy.calls.clear()
    collect(config(),now=BASE+timedelta(days=1,minutes=30),root=control,source_factory=Spy)
    assert any(local_time(c[2])<=BASE for c in Spy.calls)
    closed=tmp_path/'closed';closed.mkdir();prepare(closed)
    gap_start=BASE-timedelta(days=2)
    record_feedback(config(),station='B',reason='设备报废',unrecoverable=True,
                    start=gap_start+timedelta(hours=2),end=gap_start+timedelta(hours=4),
                    feedback_at=BASE-timedelta(days=1),root=closed)
    Spy.calls.clear()
    collect(config(),now=BASE+timedelta(days=1,minutes=40),root=closed,source_factory=Spy)
    assert Spy.calls and all(local_time(c[2])>=BASE-timedelta(days=1) for c in Spy.calls)


def test_feedback_requires_matching_station(tmp_path):
    from app.services.station_operation_monitor.feedback import record_feedback
    with pytest.raises(ValueError,match='站点名录'):
        record_feedback(config(),station='NOPE',reason='x',unrecoverable=True,root=tmp_path)


def test_pending_rechecks_only_ongoing_offline_spans(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    absent={('B',h) for h in hours(BASE,BASE+timedelta(hours=2))}
    absent.update(('B',h) for h in hours(BASE+timedelta(hours=20),end))
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    intervals=store.pending_hours()
    # The first gap has complete data after it (recovered): re-checking it only
    # ever served backfill detection, so it leaves the rotation. Only the outage
    # that still extends past the last observed hour stays pending.
    assert len(intervals)==1
    assert local_time(intervals[0]['start'])==local_time(BASE+timedelta(hours=21))
    assert local_time(intervals[0]['end'])==local_time(BASE+timedelta(hours=24))


def test_multi_hour_partial_parameter_gap_is_tracked_and_backfilled(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    records=rows(BASE,end)
    for record in records:
        if record['timePoint'] in set(hours(BASE,BASE+timedelta(hours=5))):
            record['o3']='—'
    ingest(store,BASE,end,records)
    entry=report(store,end)['entries'][0]
    assert entry['affected_hours']==5 and entry['backfill_status']=='待回补'
    assert entry['missing_parameters']==['o3']
    ingest(store,BASE,end,rows(BASE,end),observed=end+timedelta(minutes=31))
    entry=report(store,end)['entries'][0]
    assert entry['affected_hours']==5 and entry['backfill_status']=='已回补'


def test_separate_long_gaps_remain_separate_events(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    labels=list(hours(BASE,end))
    ingest(store,BASE,end,rows(BASE,end,missing={('A',h) for h in labels[0:5]}|{('A',h) for h in labels[6:11]}))
    assert [e['affected_hours'] for e in report(store,end)['entries']]==[5,5]


def test_partial_missing_at_cutoff_does_not_claim_recovery(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    records=rows(BASE,end)
    cutoff_hours=set(hours(end-timedelta(hours=5),end))
    for record in records:
        if record['code']=='A' and record['timePoint'] in cutoff_hours:
            record['o3']='—'
    ingest(store,BASE,end,records)
    entry=report(store,end)['entries'][0]
    assert not entry['recovered'] and entry['missing_parameters']==['o3']


@pytest.mark.asyncio
async def test_report_tool_publishes_word_and_pdf_to_session_catalog(tmp_path, monkeypatch):
    from app.services.station_operation_monitor import runtime, word
    from app.tools.analysis.station_operation_daily.tool import StationOperationDailyTool
    from app.agent.resources.runtime import persist_tool_result_resources
    from app.agent.resources.resource_service import SessionResourceService
    from app.scheduled_tasks.result_extraction import declared_file_paths
    doc=tmp_path/'日报.docx';doc.write_bytes(b'word')
    pdf=tmp_path/'日报.pdf';pdf.write_bytes(b'%PDF-1.4')
    monkeypatch.setattr(runtime,'generate_report',lambda **kwargs:{'success':True,'report_ready':True,'word_path':str(doc)})
    monkeypatch.setattr(word,'render_pdf_preview',lambda path:pdf)
    result=await StationOperationDailyTool().execute()
    assert len(result['resources'])==2 and 'resources' not in result['data']
    service=SessionResourceService.in_memory()
    await persist_tool_result_resources(service,'report-session','report-run',
        {'type':'tool_result','data':{'tool_name':'station_operation_daily','result':result}},turn_sequence=1)
    resources=(await service.list_resources('report-session')).resources
    assert len(resources)==2
    primary=next(r for r in resources if r.relation=='primary')
    preview=next(r for r in resources if r.relation=='preview')
    assert preview.parent_resource_id==primary.resource_id
    assert preview.renderer=='pdf' and 'preview' in preview.capabilities
    assert declared_file_paths(result)==[str(doc),str(pdf)]


def test_scheduled_result_extracts_declared_document_locators():
    from app.scheduled_tasks.models.execution import TaskExecution
    from app.scheduled_tasks.result_extraction import extract_task_result
    execution=TaskExecution(execution_id='e',task_id='t',task_name='日报',status='success',total_steps=1)
    result=extract_task_result(execution,None,{'tool_calls':[{'tool':'station_operation_daily','result':{
        'resources':[{'locator':{'path':'backend/reports/日报.docx'}},
                     {'locator':{'path':'backend/reports/日报.pdf'}}]}}]})
    assert result.document_paths==['backend/reports/日报.docx','backend/reports/日报.pdf']


def test_recovered_event_exits_at_24_hour_window_boundary(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1)
    gap_end=BASE+timedelta(hours=5)
    absent={('A',h) for h in hours(BASE,gap_end)}
    ingest(store,BASE,end,rows(BASE,end,missing=absent))
    assert report(store,gap_end+timedelta(hours=23))['entries']
    assert report(store,gap_end+timedelta(hours=24))['entries']==[]


def test_new_outage_does_not_reopen_old_recovered_event(tmp_path):
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    first_end=BASE+timedelta(days=1)
    absent={('A',h) for h in hours(BASE,BASE+timedelta(hours=5))}
    ingest(store,BASE,first_end,rows(BASE,first_end,missing=absent))
    end=first_end+timedelta(days=1)
    absent={('A',h) for h in hours(end-timedelta(hours=5),end)}
    ingest(store,first_end,end,rows(first_end,end,missing=absent))
    entries=report(store,end)['entries']
    assert len(entries)==1 and not entries[0]['recovered']
    assert local_time(entries[0]['offline_start'])==end-timedelta(hours=5)


def test_fixed_template_rejects_extra_content(tmp_path):
    import json
    from docx import Document
    from docx.shared import Pt
    from app.services.station_operation_monitor.word import write_word, validate_template_format
    source = _daily_template(tmp_path)
    store=MonitorStore(tmp_path/'state.sqlite');store.seed(ROSTER)
    end=BASE+timedelta(days=1);ingest(store,BASE,end,rows(BASE,end))
    data=report(store,end);output=tmp_path/'日报.docx'
    write_word(data,{'template_path':str(source),'prepared_by':'其他编制','reviewed_by':'其他校核'},output)
    assert json.loads(output.with_suffix('.format.json').read_text())['checks']['fixed_content']
    document=Document(output)
    assert document.paragraphs[10].text == '编制：许凡                              校核：张涛'
    assert document.paragraphs[11].text == '主送：省生态环境厅监测处                抄送：各驻市站'
    assert not any(p.text=='数据核查说明' for p in document.paragraphs)
    assert document.paragraphs[7].text == '过去24小时内，无国控站点出现过离线超4小时的情况。'
    assert document.paragraphs[9].text == '过去24小时内，无省控站点出现过离线超4小时的情况。'
    document.paragraphs[7].text = '过去24小时内，无国控站点出现数据缺失的情况。'
    document.save(output)
    with pytest.raises(ValueError, match='section_wording'):
        validate_template_format(source,output,data=data)
    write_word(data,{'template_path':str(source)},output)
    document=Document(output)
    document.paragraphs[10].text = '编制：其他人员 校核：其他人员'
    document.save(output)
    with pytest.raises(ValueError,match='fixed_signatures_recipients'):
        validate_template_format(source,output,data=data)
    write_word(data,{'template_path':str(source)},output)
    document=Document(output)
    document.add_paragraph('擅自添加的总结');document.save(output)
    with pytest.raises(ValueError,match='fixed_content'):
        validate_template_format(source,output,data=data)


@pytest.mark.asyncio
async def test_fixed_workflow_rejects_unapproved_arguments():
    from types import SimpleNamespace
    from app.services.station_operation_monitor.workflow import run_daily_report
    task=SimpleNamespace(workflow_args={'extra_paragraph':'自由发挥'})
    with pytest.raises(ValueError):
        await run_daily_report(task,SimpleNamespace())


def _daily_template(tmp_path):
    from docx import Document
    from docx.shared import Pt
    template=Document()
    labels=['广东省空气质量国控和省控站点','运行监控日报','日期','单位','分隔线','统计说明',
            '国控站点','国控摘要','事件一','事件二','省控站点','省控摘要','编制：许凡                              校核：张涛',
            '主送：省生态环境厅监测处                抄送：各驻市站',
            '','附件：','国控站点离线记录表','','省控站点离线记录表','']
    for text in labels:
        p=template.add_paragraph();p.paragraph_format.space_before=Pt(0);p.add_run(text).font.size=Pt(16)
        for _ in range(5):p.add_run('').font.size=Pt(16)
    for kind in ('国控','省控'):
        table=template.add_table(rows=2,cols=9)
        for cell in table.rows[0].cells:cell.text='表头'
        for cell in table.rows[1].cells:
            cell.text='示例';cell.paragraphs[0].runs[0].font.size=Pt(9)
        if kind == '省控':table.rows[1].cells[3].paragraphs[0].runs[0].font.size=None
    source=tmp_path/'template.docx';template.save(source)
    return source


def test_report_text_and_table_match_business_template(tmp_path):
    from docx import Document
    from app.services.station_operation_monitor.word import write_word
    source = _daily_template(tmp_path)
    common = dict(city='揭阳', name='产业园龙尾', type='省控', recovered=False,
                  offline_start='2026-10-09T17:00:00+08:00', offline_end='2026-10-12T10:00:00+08:00',
                  affected_hours=65, carry_over=True, backfill_status='待回补',
                  backfill_sources=['内部接口名'], missing_parameters=['co', 'nO2', 'o3', 'pM10', 'pM2_5', 'sO2'],
                  feedback={'feedback_by':'核实人', 'feedback_at':'2026-10-12T09:00:00+08:00', 'reason':'区域网络基站异常'})
    unresolved = dict(common, issue_reason='区域网络基站异常', data_validity_impact='否')
    recovered = dict(common, city='东莞', name='南城元美', affected_hours=5, recovered=True,
                     offline_start='2026-10-12T02:00:00+08:00', offline_end='2026-10-12T07:00:00+08:00',
                     issue_reason='数采软件故障', backfill_status='已回补', data_validity_impact='否')
    unknown = dict(common, city='茂名', name='茂石化七小', type='国控', affected_hours=14,
                   offline_start='2026-10-11T20:00:00+08:00')
    data = dict(window_end='2026-10-12T10:00:00+08:00', station_counts={'国控':122, '省控':163},
                entries=[unknown, unresolved, recovered])
    output=tmp_path/'日报.docx'
    write_word(data, {'template_path':str(source)}, output)
    d=Document(output)
    text='\n'.join(p.text for p in d.paragraphs)
    assert '过去24小时内，有1个国控站点出现过离线超4小时的情况，经核实，情况如下：' in text
    assert '过去24小时内，有2个省控站点出现过离线超4小时的情况，经核实，情况如下：' in text
    assert '其余国控和省控站点运行正常。' in d.paragraphs[5].text
    assert '具体如下' not in text
    assert '揭阳市产业园龙尾站：经核实为区域网络基站异常，累计影响数据时长65小时，正在持续跟进。' in text
    assert '东莞市南城元美站：经核实为数采软件故障，累计影响数据时长5小时，现已恢复正常，数据已回补。' in text
    assert '茂名市茂石化七小站：累计影响数据时长14小时，正在持续跟进。' in text
    assert [c.text for c in d.tables[1].rows[1].cells] == [
        '1','揭阳','产业园龙尾','10月9日17时-12日10时','否','65','否','区域网络基站异常','持续跟进']
    assert [c.text for c in d.tables[1].rows[2].cells] == [
        '2','东莞','南城元美','10月12日2时-7时','是','5','否','数采软件故障','数据已回补']
    assert [c.text for c in d.tables[0].rows[1].cells][-3:] == ['/','/','持续跟进']
    assert d.tables[1].rows[1].cells[3].paragraphs[0].runs[0].font.size.pt == 9
    all_text=text+'\n'+'\n'.join(c.text for table in d.tables for row in table.rows for c in row.cells)
    for forbidden in ('缺失项目', 'CO', '内部接口名', '历史未结跟进', '核实人', '关闭跟进', '原因待核实'):
        assert forbidden not in all_text


@pytest.mark.parametrize('recovered,status,note', [
    (True, '待回补', '数据待回补'), (True, '已回补', '数据已回补'),
    (False, '无法回补', '数据无法回补'), (False, '待回补', '持续跟进')])
def test_business_backfill_notes_keep_recovery_separate(recovered, status, note):
    from app.services.station_operation_monitor.word import incident_text, incident_table_values
    entry=dict(city='测试', name='测试', recovered=recovered, backfill_status=status,
               affected_hours=6, offline_start=BASE.isoformat(), offline_end=(BASE+timedelta(hours=6)).isoformat())
    values=incident_table_values(entry, 1)
    assert values[4] == ('是' if recovered else '否')
    assert values[-1] == note
    if recovered or status == '无法回补':
        assert note in incident_text(entry)


@pytest.mark.parametrize('kinds,expected', [
    (['省控'], '有1个省控站点出现过离线超4小时的情况，其余国控和省控站点运行正常。'),
    (['国控'], '有1个国控站点出现过离线超4小时的情况，其余国控和省控站点运行正常。'),
    (['国控','省控'], '有1个国控站点和1个省控站点出现过离线超4小时的情况，其余国控和省控站点运行正常。'),
    ([], '国控和省控站点运行正常。')])
def test_overview_uses_exact_template_wording(kinds, expected):
    from app.services.station_operation_monitor.word import overview_summary
    entries=[dict(type=kind, city='测试', name=kind) for kind in kinds]
    # Repeated incidents at one station must not inflate the station count.
    data=dict(window_end='2026-10-19T10:00:00+08:00', report_ready=True, entries=entries+entries)
    prefix='通过广东省空气质量监测管理与发布系统检查、核对，截至10月19日10时，在过去24小时内，'
    assert overview_summary(data) == prefix + expected
    data['report_ready']=False
    assert '运行正常' not in overview_summary(data)


def test_fixed_note_is_11pt_in_first_page_footer_only(tmp_path):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from app.services.station_operation_monitor.word import FIRST_PAGE_NOTE, write_word, validate_template_format
    source=_daily_template(tmp_path)
    template=Document(source)
    page_field=OxmlElement('w:fldSimple');page_field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr','PAGE')
    template.sections[0].footer.paragraphs[0]._p.append(page_field)
    template.save(source)
    data=dict(window_end='2026-10-20T10:00:00+08:00', station_counts={'国控':122,'省控':163}, entries=[])
    output=tmp_path/'fixed_note.docx'
    write_word(data,{'template_path':str(source)},output)
    document=Document(output)
    section=document.sections[0]
    assert section.different_first_page_header_footer
    note=next(p for p in section.first_page_footer.paragraphs if p.text==FIRST_PAGE_NOTE)
    assert all(r.font.size==Pt(11) for r in note.runs)
    assert 'PAGE' in section.first_page_footer._element.xml
    assert 'PAGE' in section.footer._element.xml
    assert FIRST_PAGE_NOTE not in '\n'.join(p.text for p in section.footer.paragraphs)
    assert FIRST_PAGE_NOTE not in '\n'.join(p.text for p in document.paragraphs)
    note.runs[0].font.size=Pt(12)
    document.save(output)
    with pytest.raises(ValueError, match='first_page_note'):
        validate_template_format(source,output,data=data)
